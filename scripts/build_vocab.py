#!/usr/bin/env python3
"""
Rebuilds data/vocab.json from the "Glossary chapter *.docx" Word files.

Usage:
    python3 scripts/build_vocab.py "/path/to/word lists and games" [--out data/vocab.json]

Each source .docx is expected to contain one table with a header row
("FIN", "ENG") followed by one Finnish/English word pair per row -- this is
the format used by the Metropolia Finnish A1.2 glossary handouts.

Re-run this whenever a new "Glossary chapter N.docx" file is added to the
course folder (e.g. chapter 6, 7, ...): it picks up any file matching
"Glossary chapter*.docx" automatically, so no code changes are needed.

Contextual "hint" text (shown by the Write-mode Hint button) is stored
separately in data/hints.json, keyed by word id, and merged in here -- it is
NOT regenerated from the docx files, since it takes human/editorial judgment.
Any word id with no entry in hints.json simply has no hint (the Hint button
falls back to a generic behavior for it). After adding a new chapter, add
hint entries for its new ids to hints.json if you want hints for them.
"""
import argparse
import glob
import json
import os
import re
import sys

try:
    import docx
except ImportError:
    sys.exit(
        "Missing dependency 'python-docx'. Install it with:\n"
        "    pip install python-docx --break-system-packages\n"
        "(or just: pip install python-docx)"
    )

CHAPTER_LABEL_OVERRIDES = {
    "1-2": "Chapters 1-2",
}


def chapter_label_from_filename(filename: str) -> str:
    """'Glossary chapter 3.docx' -> 'Chapter 3'; 'Glossary chapter 1-2.docx' -> 'Chapters 1-2'."""
    m = re.search(r"chapter\s+([0-9-]+)", filename, re.IGNORECASE)
    if not m:
        return os.path.splitext(filename)[0]
    key = m.group(1)
    if key in CHAPTER_LABEL_OVERRIDES:
        return CHAPTER_LABEL_OVERRIDES[key]
    return f"Chapter {key}"


def chapter_sort_key(label: str):
    nums = [int(n) for n in re.findall(r"\d+", label)]
    return nums[0] if nums else 999


def extract_pairs(path: str):
    d = docx.Document(path)
    if not d.tables:
        return []
    table = d.tables[0]
    pairs = []
    for row in table.rows[1:]:  # skip header ("FIN", "ENG")
        cells = [c.text.strip() for c in row.cells]
        if len(cells) < 2:
            continue
        fi = re.sub(r"\s+", " ", cells[0]).strip()
        en = re.sub(r"\s+", " ", cells[1]).strip()
        # The glossaries mark verb types with a trailing number ("soittaa 1",
        # "haluta 4"). That's a grammar note, not part of the word, so drop it
        # -- otherwise Write mode would expect learners to type the digit.
        verb_type = None
        m = re.search(r"\s+(\d+)$", fi)
        if m:
            verb_type = int(m.group(1))
            fi = fi[: m.start()].strip()
        if not fi and not en:
            continue
        pairs.append((fi, en, verb_type))
    return pairs


def load_hints(hints_path: str):
    if not os.path.exists(hints_path):
        return {}
    with open(hints_path, "r", encoding="utf-8") as f:
        raw = json.load(f)
    return {k: v for k, v in raw.items() if not k.startswith("_")}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source_dir", help="Folder containing 'Glossary chapter *.docx' files")
    ap.add_argument("--out", default="data/vocab.json", help="Output JSON path")
    ap.add_argument(
        "--hints",
        default=None,
        help="Path to hints.json (default: hints.json next to --out)",
    )
    args = ap.parse_args()

    pattern = os.path.join(args.source_dir, "Glossary chapter*.docx")
    files = sorted(glob.glob(pattern))
    if not files:
        sys.exit(f"No files matching 'Glossary chapter*.docx' found in {args.source_dir}")

    by_file = {}
    for path in files:
        fname = os.path.basename(path)
        label = chapter_label_from_filename(fname)
        by_file.setdefault(label, []).extend(extract_pairs(path))

    ordered_labels = sorted(by_file.keys(), key=chapter_sort_key)

    out_path = args.out
    hints_path = args.hints or os.path.join(os.path.dirname(out_path) or ".", "hints.json")
    hints = load_hints(hints_path)

    words = []
    uid = 1
    missing_hints = []
    for label in ordered_labels:
        for fi, en, verb_type in by_file[label]:
            word = {"id": uid, "chapter": label, "fi": fi, "en": en}
            if verb_type:
                word["verbType"] = verb_type
            hint = hints.get(str(uid))
            if hint:
                word["hint"] = hint
            else:
                missing_hints.append(uid)
            words.append(word)
            uid += 1

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(
            {"chapters": ordered_labels, "words": words},
            f,
            ensure_ascii=False,
            indent=2,
        )

    print(f"Wrote {len(words)} words across {len(ordered_labels)} chapters -> {out_path}")
    for label in ordered_labels:
        print(f"  {label}: {len(by_file[label])} words")
    if missing_hints:
        print(
            f"\n{len(missing_hints)} word(s) have no hint yet in {hints_path} "
            f"(ids: {missing_hints[0]}-{missing_hints[-1]}). "
            "The Write-mode Hint button falls back to a generic hint for these."
        )


if __name__ == "__main__":
    main()
